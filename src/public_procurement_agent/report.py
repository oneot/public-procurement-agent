"""수집 결과 리포트 출력 (CSV · 마크다운 요약 · 엑셀)."""

from __future__ import annotations

import csv
import logging
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from public_procurement_agent.collector import CollectionResult
from public_procurement_agent.config import AppConfig
from public_procurement_agent.formatting import format_won
from public_procurement_agent.models import CSV_COLUMNS, BidNotice

log = logging.getLogger(__name__)

CSV_PREFIX = "초중고교육청_입찰공고_"
SUMMARY_PREFIX = "초중고교육청_요약_"

_XLSX_WIDTHS = (12, 8, 16, 6, 50, 26, 22, 12, 15, 18, 18, 18, 10, 16, 6, 40)
_EOK = 100_000_000


@dataclass(frozen=True, slots=True)
class ReportPaths:
    csv: Path
    summary: Path
    xlsx: Path | None  # openpyxl 미설치 시 None


def csv_path_for(output_dir: Path, day: date) -> Path:
    return output_dir / f"{CSV_PREFIX}{day:%Y%m%d}.csv"


def write_reports(
    result: CollectionResult,
    config: AppConfig,
    *,
    output_dir: Path,
    today: date,
    generated_at: datetime,
) -> ReportPaths:
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = f"{today:%Y%m%d}"
    csv_path = csv_path_for(output_dir, today)
    summary_path = output_dir / f"{SUMMARY_PREFIX}{stamp}.md"
    xlsx_path = output_dir / f"{CSV_PREFIX}{stamp}.xlsx"

    write_csv(result.notices, csv_path)
    write_summary(
        result,
        summary_path,
        topic_names=list(config.filters.topics),
        large_budget_mark=config.collect.large_budget_mark,
        generated_at=generated_at,
    )
    wrote_xlsx = write_xlsx(result.notices, xlsx_path)
    return ReportPaths(csv=csv_path, summary=summary_path, xlsx=xlsx_path if wrote_xlsx else None)


def write_csv(notices: Sequence[BidNotice], path: Path) -> None:
    """엑셀에서 바로 열 수 있도록 UTF-8 BOM 으로 저장합니다."""
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(CSV_COLUMNS)
        writer.writerows(notice.csv_row() for notice in notices)


def _large_label(mark: int) -> str:
    if mark >= _EOK and mark % _EOK == 0:
        return f"{mark // _EOK}억+"
    return f"{format_won(mark)}+"


def render_summary(
    result: CollectionResult,
    *,
    topic_names: Sequence[str],
    large_budget_mark: int,
    generated_at: datetime,
) -> str:
    notices = sorted(result.notices, key=lambda n: n.budget, reverse=True)
    by_topic = Counter(topic for notice in notices for topic in notice.topics)
    by_category = Counter(notice.category for notice in notices)
    large_count = sum(1 for notice in notices if notice.is_large)

    lines = [
        f"# 나라장터 입찰공고 — 초·중·고/교육청 {'·'.join(topic_names)} 사업 기회",
        f"- 수집 기간: {result.period.label} (등록일시 기준)",
        f"- 생성 시각: {generated_at:%Y-%m-%d %H:%M}",
        "",
        "## 한눈에 보기",
        f"- 전체 조회: **{result.scanned}건** → 교육기관: **{result.k12_count}건** "
        f"→ **초점 주제 매칭: {len(notices)}건** (최종 대상)",
        "- 주제별: "
        + ", ".join(f"{topic} {by_topic[topic]}건" for topic in topic_names if by_topic[topic]),
        "- 업무별: "
        + ", ".join(f"{category} {count}건" for category, count in by_category.items()),
        f"- 대형({_large_label(large_budget_mark)}): **{large_count}건**",
        "",
        "## 사업 기회 목록 (예산순)",
    ]
    if not notices:
        lines.append("- 조건에 맞는 공고가 없습니다.")
    for notice in notices:
        mark = "★ " if notice.is_large else ""
        lines.append(f"- {mark}`[{'/'.join(notice.topics)}]` **{notice.name}**")
        lines.append(
            f"  - {notice.customer} · {format_won(notice.budget)} · {notice.category} "
            f"· 마감 {notice.close_dt or '미정'}"
        )
        if notice.url:
            lines.append(f"  - {notice.url}")
    lines.append("")
    return "\n".join(lines)


def write_summary(
    result: CollectionResult,
    path: Path,
    *,
    topic_names: Sequence[str],
    large_budget_mark: int,
    generated_at: datetime,
) -> None:
    path.write_text(
        render_summary(
            result,
            topic_names=topic_names,
            large_budget_mark=large_budget_mark,
            generated_at=generated_at,
        ),
        encoding="utf-8",
    )


def write_xlsx(notices: Sequence[BidNotice], path: Path) -> bool:
    """엑셀 리포트를 저장합니다. ``openpyxl`` 이 없으면 건너뛰고 ``False``."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.worksheet import Worksheet
    except ImportError:
        log.info("openpyxl 미설치 → xlsx 생략 (설치: uv sync --extra xlsx)")
        return False

    workbook = Workbook()
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):  # pragma: no cover - 새 통합문서는 항상 워크시트를 가짐
        raise TypeError("openpyxl 활성 시트가 워크시트가 아닙니다")
    sheet.title = "사업기회"
    sheet.append(list(CSV_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")
    for notice in sorted(notices, key=lambda n: n.budget, reverse=True):
        sheet.append(notice.csv_row())
        for cell in sheet[sheet.max_row]:
            # '=' 로 시작하는 공고명 등이 수식으로 해석되지 않도록 문자열로 고정합니다.
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    for index, width in enumerate(_XLSX_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)
    return True
